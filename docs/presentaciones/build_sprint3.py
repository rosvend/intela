#!/usr/bin/env python3
"""Regenerate Intela Sprint 3 deck in the Avances-Grupo2 visual system."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os
import copy

ASSETS = os.path.join(os.path.dirname(__file__), "_assets")
OUT = os.path.join(os.path.dirname(__file__), "Intela-Sprint-3.pptx")
OUT_DL = "/Users/emanuel.acevedo/Downloads/Intela-Sprint-3.pptx"

# Palette from Intela-Avances-Grupo2
MAGENTA = RGBColor(0xC9, 0x18, 0x6B)
NEAR_BLACK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
UPB_GRAY = RGBColor(0x4A, 0x4A, 0x4A)
CARD = RGBColor(0xF4, 0xF4, 0xF6)
CARD_HOT = RGBColor(0xF7, 0xE4, 0xEE)
PILL_GRAY = RGBColor(0xD9, 0xD9, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DONE = RGBColor(0x2E, 0x7D, 0x4F)
WIP = RGBColor(0xB5, 0x6A, 0x1B)
OPEN = RGBColor(0x9B, 0x2C, 0x2C)

UPB_BARS = os.path.join(ASSETS, "upb_bars.png")
UPB_CREST = os.path.join(ASSETS, "upb_crest.png")

# Screenshots reused from prior Sprint-3 deck
IMG = {i: os.path.join(ASSETS, f"s3_image{i}.png") for i in range(1, 10)}

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

ML = Emu(777240)  # left margin matching Avances
MR = Emu(777240)
CONTENT_W = W - ML - MR
TOTAL = 15


def set_run(run, size, color, bold=False, name="Calibri"):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name


def add_text(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold) or str."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign])
    except Exception:
        pass
    first = True
    for item in lines:
        if isinstance(item, str):
            item = (item, 14, NEAR_BLACK, False)
        text, size, color, bold = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = text
        set_run(run, size, color, bold)
    return box


def add_title_mixed(slide, left, top, width, height, parts, size=30):
    """parts: list of (text, color, bold)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for text, color, bold in parts:
        run = p.add_run()
        run.text = text
        set_run(run, size, color, bold)
    return box


def rect(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def card(slide, left, top, width, height, fill=CARD, stripe=True):
    body = rect(slide, left, top, width, height, fill)
    if stripe:
        rect(slide, left, top, width, Emu(63720), MAGENTA)
    return body


def upb_header(slide):
    slide.shapes.add_picture(UPB_BARS, Emu(9674280), Emu(457200), width=Emu(1737000), height=Emu(580320))


def upb_footer(slide, page):
    add_text(
        slide,
        ML,
        Emu(6382440),
        Emu(7314840),
        Emu(273960),
        [(f"Intela  ·  Grupo 2  ·  Sprint 3  ·  UPB  ·  {page}", 10, MUTED, False)],
    )
    slide.shapes.add_picture(UPB_CREST, Emu(9802440), Emu(6126480), width=Emu(525960), height=Emu(566640))
    add_text(
        slide,
        Emu(10420200),
        Emu(6144840),
        Emu(1279800),
        Emu(566640),
        [
            ("Universidad", 9, UPB_GRAY, True),
            ("Pontificia", 9, UPB_GRAY, True),
            ("Bolivariana", 9, UPB_GRAY, True),
        ],
        valign=MSO_ANCHOR.MIDDLE,
    )


def header_block(slide, eyebrow, title_parts, page):
    """Standard content-slide chrome."""
    upb_header(slide)
    add_text(slide, ML, Emu(567000), CONTENT_W - Emu(2000000), Emu(273960), [(eyebrow.upper(), 12, MUTED, True)])
    if isinstance(title_parts, str):
        title_parts = [(title_parts, NEAR_BLACK, True)]
    add_title_mixed(slide, ML, Emu(868680), CONTENT_W - Emu(2000000), Emu(1051200), title_parts, size=30)
    # accent under title
    rect(slide, ML, Emu(1828800), Emu(5120280), Emu(100080), MAGENTA)
    upb_footer(slide, page)


def new_slide():
    return prs.slides.add_slide(BLANK)


# ---------------------------------------------------------------------------
# SLIDE 1 — Title (Avances title style)
# ---------------------------------------------------------------------------
s = new_slide()
upb_header(s)
# magenta bar + Intela
rect(s, ML, Emu(2100000), Emu(820000), Emu(100080), MAGENTA)
add_text(s, Emu(1700000), Emu(1900000), Emu(5000000), Emu(700000), [("Intela", 54, MAGENTA, True)])
add_text(
    s,
    ML,
    Emu(2700000),
    Emu(7000000),
    Emu(500000),
    [("Sprint 3 · Del reporte de un canal al reparto de un guionista", 22, NEAR_BLACK, False)],
)
add_text(
    s,
    ML,
    Emu(3300000),
    Emu(7000000),
    Emu(600000),
    [
        ("Sistema de reconocimiento de obras y distribución de ingresos", 14, MUTED, False),
        ("por propiedad intelectual para REDES SGC", 14, MUTED, False),
    ],
)
add_text(
    s,
    ML,
    Emu(4800000),
    Emu(5500000),
    Emu(1200000),
    [
        ("Emanuel Acevedo", 14, NEAR_BLACK, False),
        ("Roy Sandoval", 14, NEAR_BLACK, False),
        ("Miguel Legarda", 14, NEAR_BLACK, False),
        ("Santiago Mendoza", 14, NEAR_BLACK, False),
    ],
)
add_text(
    s,
    ML,
    Emu(6300000),
    Emu(7000000),
    Emu(300000),
    [("Proyecto Aplicado en TIC II  ·  Sprint 3  ·  Septiembre de 2026", 11, MUTED, False)],
)
# catalog screenshot
s.shapes.add_picture(IMG[1], Emu(7800000), Emu(2500000), width=Emu(3800000), height=Emu(2375000))
add_text(
    s,
    Emu(7800000),
    Emu(4950000),
    Emu(3800000),
    Emu(300000),
    [("El catálogo de obras, ya en producción", 11, MUTED, False)],
    align=PP_ALIGN.CENTER,
)

# ---------------------------------------------------------------------------
# SLIDE 2 — Punto de partida
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "El punto de partida",
    [("Un problema de dinero ajeno ", MAGENTA, True), ("y de papeles", NEAR_BLACK, True)],
    1,
)
card(s, ML, Emu(2100000), CONTENT_W, Emu(900000), CARD_HOT)
add_text(
    s,
    ML + Emu(200000),
    Emu(2250000),
    CONTENT_W - Emu(400000),
    Emu(700000),
    [
        (
            "Un guionista cobra cuando su obra se emite. Hoy, calcular cuánto le toca es un trabajo manual sobre miles de filas de parrilla, para cada canal y cada periodo.",
            15,
            NEAR_BLACK,
            False,
        )
    ],
)
cols = [
    ("Lo que dice la ley", "La Ley 1835 de 2017 (Ley Pepe Sánchez) creó el derecho de remuneración de los guionistas y libretistas por la comunicación pública de sus obras."),
    ("Quién paga y cuánto", "Canales, plataformas, salas de cine, hoteles y transporte. Cada uno paga una tarifa concertada y distinta: el convenio firmado manda."),
    ("La exigencia de fondo", "La auditoría puede llegar en cualquier momento y los registros se conservan diez años. Cada cifra debe explicarse hasta su origen."),
]
cw = Emu(3428640)
gap = Emu(180000)
for i, (t, b) in enumerate(cols):
    left = ML + i * (cw + gap)
    card(s, left, Emu(3200000), cw, Emu(2200000))
    add_text(s, left + Emu(180000), Emu(3400000), cw - Emu(360000), Emu(500000), [(t, 14, NEAR_BLACK, True)])
    add_text(s, left + Emu(180000), Emu(3950000), cw - Emu(360000), Emu(1200000), [(b, 12, MUTED, False)])
add_text(
    s,
    ML,
    Emu(5550000),
    CONTENT_W,
    Emu(400000),
    [("El sistema no inventa cifras: si falta un dato, se detiene y lo dice.", 13, MAGENTA, True)],
    align=PP_ALIGN.CENTER,
)

# ---------------------------------------------------------------------------
# SLIDE 3 — La idea
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "La idea que cambia todo",
    [("El dinero no llega ", MAGENTA, True), ("por fila", NEAR_BLACK, True)],
    2,
)
s.shapes.add_picture(IMG[2], ML, Emu(2100000), width=CONTENT_W, height=Emu(2500000))
card(s, ML, Emu(4750000), CONTENT_W, Emu(1400000), CARD_HOT)
add_text(
    s,
    ML + Emu(200000),
    Emu(4850000),
    CONTENT_W - Emu(400000),
    Emu(400000),
    [("La analogía que lo explica en diez segundos", 14, NEAR_BLACK, True)],
)
add_text(
    s,
    ML + Emu(200000),
    Emu(5250000),
    CONTENT_W - Emu(400000),
    Emu(800000),
    [
        (
            "Imagina diez locales de un centro comercial que pagan entre todos una cuota por la música que suena en las zonas comunes. La cuota no dice qué canción sonó: eso lo dice el reporte de uso. Primero se reparte entre las canciones según cuánto sonaron; después, cada canción paga a sus autores según lo que cada uno declaró.",
            12,
            MUTED,
            False,
        )
    ],
)

# ---------------------------------------------------------------------------
# SLIDE 4 — Cálculo
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "El cálculo, con las cifras del reglamento",
    [("Canal Z: de la bolsa ", MAGENTA, True), ("al pago de cada obra", NEAR_BLACK, True)],
    3,
)
card(s, ML, Emu(2100000), CONTENT_W, Emu(900000))
add_text(
    s,
    ML + Emu(200000),
    Emu(2200000),
    CONTENT_W - Emu(400000),
    Emu(350000),
    [("Puntos = ponderación × duración × rating × emisiones", 14, NEAR_BLACK, True)],
)
add_text(
    s,
    ML + Emu(200000),
    Emu(2550000),
    CONTENT_W - Emu(400000),
    Emu(350000),
    [("Valor del punto = total del canal ÷ total de puntos del canal", 14, NEAR_BLACK, True)],
)
add_text(
    s,
    ML,
    Emu(3050000),
    CONTENT_W,
    Emu(350000),
    [
        (
            "La duración artística es el 80 % de la reportada y la hora de televisión se computa como 48 minutos. Las dos reglas no se encadenan.",
            12,
            MUTED,
            False,
        )
    ],
)
metrics = [
    ("1.000.000", "la bolsa del canal Z, tras deducciones"),
    ("139,1", "el valor del punto ese periodo"),
    ("$219.024", "Película X · 1 emisión"),
    ("$780.976", "Serie Y · 10 emisiones"),
]
mw = Emu(2500000)
mg = Emu(160000)
for i, (n, lab) in enumerate(metrics):
    left = ML + i * (mw + mg)
    card(s, left, Emu(3500000), mw, Emu(1600000), CARD_HOT if i < 2 else CARD)
    add_text(s, left + Emu(120000), Emu(3700000), mw - Emu(240000), Emu(600000), [(n, 26, MAGENTA, True)], align=PP_ALIGN.CENTER)
    add_text(s, left + Emu(120000), Emu(4400000), mw - Emu(240000), Emu(500000), [(lab, 11, MUTED, False)], align=PP_ALIGN.CENTER)
add_text(
    s,
    ML,
    Emu(5300000),
    CONTENT_W,
    Emu(700000),
    [
        ("Es un caso de prueba del sistema.", 12, NEAR_BLACK, True),
        ("Si el motor no reproduce estas cifras del reglamento, la compuerta de integración falla.", 12, MUTED, False),
    ],
)

# ---------------------------------------------------------------------------
# SLIDE 5 — Entregado
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Lo que se entregó",
    [("Sprint 3, ", MAGENTA, True), ("en un vistazo", NEAR_BLACK, True)],
    4,
)
stats = [
    ("16 de 17", "issues entregadas"),
    ("6", "frentes de trabajo"),
    ("21 sep", "cierre del sprint"),
]
sw = Emu(3428640)
sg = Emu(180000)
for i, (n, lab) in enumerate(stats):
    left = ML + i * (sw + sg)
    card(s, left, Emu(2100000), sw, Emu(1100000), CARD_HOT)
    add_text(s, left + Emu(120000), Emu(2250000), sw - Emu(240000), Emu(500000), [(n, 28, MAGENTA, True)], align=PP_ALIGN.CENTER)
    add_text(s, left + Emu(120000), Emu(2800000), sw - Emu(240000), Emu(300000), [(lab, 13, MUTED, False)], align=PP_ALIGN.CENTER)
s.shapes.add_picture(IMG[3], ML, Emu(3400000), width=CONTENT_W, height=Emu(2500000))

# ---------------------------------------------------------------------------
# SLIDE 6 — Frente 1 Bolsa
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Frente 1",
    [("La bolsa: ", MAGENTA, True), ("cuánto se cobró", NEAR_BLACK, True)],
    5,
)
add_text(
    s,
    ML,
    Emu(2050000),
    CONTENT_W,
    Emu(500000),
    [
        (
            "Antes de repartir hay que saber cuánto hay. El sistema recibe el importe ya cobrado por cada usuario y lo deja trazable hasta su factura.",
            14,
            MUTED,
            False,
        )
    ],
)
items = [
    ("Una bolsa por usuario y periodo", "El canal, la plataforma o el hotel que pagó, con su periodo y su circuito."),
    ("Nacional e internacional, separados", "Se recaudan y se invierten por separado, y producen dos repartos distintos."),
    ("El pagador es una entidad", "No un texto libre: usuarios de recaudo con su identificación."),
    ("La procedencia queda registrada", "convenio · tarifa · factura — responde de dónde salió el dinero."),
]
iw = Emu(5200000)
ih = Emu(1100000)
for i, (t, b) in enumerate(items):
    row, col = divmod(i, 2)
    left = ML + col * (iw + Emu(200000))
    top = Emu(2650000) + row * (ih + Emu(150000))
    card(s, left, top, iw, ih)
    add_text(s, left + Emu(180000), top + Emu(200000), iw - Emu(360000), Emu(350000), [(t, 14, NEAR_BLACK, True)])
    add_text(s, left + Emu(180000), top + Emu(550000), iw - Emu(360000), Emu(450000), [(b, 12, MUTED, False)])
add_text(
    s,
    ML,
    Emu(5550000),
    CONTENT_W,
    Emu(400000),
    [("GET /recaudo · GET /bolsas · GET /bolsas/{id}  — expuesto en el contrato de la API y cubierto por pruebas.", 12, MAGENTA, True)],
)

# ---------------------------------------------------------------------------
# SLIDE 7 — Frente 2 Ingesta
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Frente 2",
    [("Ingesta: tres fuentes, ", MAGENTA, True), ("un mismo esquema", NEAR_BLACK, True)],
    6,
)
s.shapes.add_picture(IMG[4], ML, Emu(2050000), width=Emu(6200000), height=Emu(2800000))
points = [
    ("Tres fuentes ya integradas", "Televisión, plataforma y cine, en Excel, CSV y JSON."),
    ("Añadir una fuente no toca el motor", "Cada fuente es un mapa de columnas declarativo."),
    ("Nada se descarta en silencio", "Cada rechazo queda con su motivo y su campo."),
    ("10.000 registros, menos de 5 minutos", "Verificado en la integración continua."),
    ("Las reglas del reglamento, al cargar", "Duración artística 80 % y hora de TV de 48 minutos."),
]
pw = Emu(4200000)
ph = Emu(700000)
for i, (t, b) in enumerate(points):
    top = Emu(2050000) + i * (ph + Emu(40000))
    left = Emu(7600000)
    card(s, left, top, pw, ph)
    add_text(s, left + Emu(140000), top + Emu(100000), pw - Emu(280000), Emu(280000), [(t, 12, NEAR_BLACK, True)])
    add_text(s, left + Emu(140000), top + Emu(360000), pw - Emu(280000), Emu(280000), [(b, 11, MUTED, False)])

# ---------------------------------------------------------------------------
# SLIDE 8 — Frente 3 Catálogo
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Frente 3",
    [("Catálogo de obras: ", MAGENTA, True), ("la ficha de cada obra", NEAR_BLACK, True)],
    7,
)
s.shapes.add_picture(IMG[5], ML, Emu(2050000), width=Emu(6200000), height=Emu(3800000))
bullets = [
    "Búsqueda por título, género, IPI y año.",
    "Filtros combinables en una misma consulta.",
    "Paginación sobre una instantánea coherente.",
    "El buscador usa el índice del catálogo.",
    "Cada obra queda ligada a sus usos.",
]
pw = Emu(4200000)
card(s, Emu(7600000), Emu(2050000), pw, Emu(2800000))
add_text(
    s,
    Emu(7780000),
    Emu(2200000),
    pw - Emu(360000),
    Emu(2400000),
    [(f"— {b}", 13, NEAR_BLACK, False) for b in bullets],
)
card(s, Emu(7600000), Emu(5000000), pw, Emu(1000000), CARD_HOT)
add_text(
    s,
    Emu(7780000),
    Emu(5100000),
    pw - Emu(360000),
    Emu(350000),
    [("Cada cambio deja asiento", 13, NEAR_BLACK, True)],
)
add_text(
    s,
    Emu(7780000),
    Emu(5450000),
    pw - Emu(360000),
    Emu(450000),
    [("El alta y la corrección se registran en una bitácora inmutable.", 12, MUTED, False)],
)

# ---------------------------------------------------------------------------
# SLIDE 9 — Frente 4 Declaración
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Frente 4",
    [("Declaración de Obra: ", MAGENTA, True), ("los porcentajes", NEAR_BLACK, True)],
    8,
)
s.shapes.add_picture(IMG[6], ML, Emu(2050000), width=Emu(5800000), height=Emu(4000000))
decls = [
    ("El único origen válido", "Solo la declaración del autor: ni reportes ni contratos."),
    ("Versionada por vigencia", "Una versión no se borra: se cierra y se abre otra."),
    ("Invariante del 100 %", "Si no suman 100, se retiene el total de esa obra."),
    ("Aviso antes de guardar", "El editor muestra la suma mientras se escribe."),
]
dw = Emu(4600000)
dh = Emu(900000)
for i, (t, b) in enumerate(decls):
    top = Emu(2050000) + i * (dh + Emu(80000))
    card(s, Emu(7200000), top, dw, dh)
    add_text(s, Emu(7380000), top + Emu(150000), dw - Emu(360000), Emu(300000), [(t, 13, NEAR_BLACK, True)])
    add_text(s, Emu(7380000), top + Emu(450000), dw - Emu(360000), Emu(350000), [(b, 12, MUTED, False)])

# ---------------------------------------------------------------------------
# SLIDE 10 — Frente 5 Identificación
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Frente 5",
    [("Identificar la obra ", MAGENTA, True), ("que el reporte nombra", NEAR_BLACK, True)],
    9,
)
s.shapes.add_picture(IMG[7], ML, Emu(2050000), width=CONTENT_W, height=Emu(2600000))
steps = [
    ("Escalón 1 · alias", "Los nombres con los que cada fuente ya conoce la obra."),
    ("Escalón 2 · identificadores", "IDA, EIDR e IMDB: la ficha internacional de la obra."),
    ("Filtro de repertorio", "Lo que no es del repertorio se excluye, y no cuenta como obra no identificada."),
]
sw = Emu(3428640)
sg = Emu(180000)
for i, (t, b) in enumerate(steps):
    left = ML + i * (sw + sg)
    card(s, left, Emu(4850000), sw, Emu(1200000), CARD_HOT if i == 2 else CARD)
    add_text(s, left + Emu(160000), Emu(5000000), sw - Emu(320000), Emu(350000), [(t, 13, NEAR_BLACK, True)])
    add_text(s, left + Emu(160000), Emu(5400000), sw - Emu(320000), Emu(500000), [(b, 12, MUTED, False)])

# ---------------------------------------------------------------------------
# SLIDE 11 — Frente 6 Motor
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Frente 6",
    [("El motor de reparto ", MAGENTA, True), ("y sus parámetros", NEAR_BLACK, True)],
    10,
)
card(s, ML, Emu(2050000), CONTENT_W, Emu(700000), CARD_HOT)
add_text(
    s,
    ML + Emu(200000),
    Emu(2200000),
    CONTENT_W - Emu(400000),
    Emu(450000),
    [
        (
            "Puntos = ponderación × duración × rating × emisiones   ·   Valor del punto = total del canal ÷ total de puntos",
            13,
            NEAR_BLACK,
            True,
        )
    ],
    align=PP_ALIGN.CENTER,
)
rules = [
    ("Una función, misma cifra siempre", "No lee la base, ni el reloj, ni la red: todo entra como argumento."),
    ("El valor del punto es por canal", "Cada canal se reparte conforme al pago que hizo ese periodo."),
    ("Parametrizado, no cableado", "Tarifas, deducciones y reglas viven como datos con vigencia."),
    ("Snapshot congelado al abrir la corrida", "Se fija la versión exacta de parámetros y archivos: un reproceso de dentro de diez años lee lo mismo."),
    ("Reproducibilidad en cada envío", "La CI reejecuta el motor y compara byte a byte contra la corrida de referencia."),
]
rw = CONTENT_W
rh = Emu(580000)
for i, (t, b) in enumerate(rules):
    top = Emu(2900000) + i * (rh + Emu(40000))
    card(s, ML, top, rw, rh)
    add_text(s, ML + Emu(200000), top + Emu(80000), Emu(4500000), Emu(400000), [(t, 13, NEAR_BLACK, True)])
    add_text(s, ML + Emu(4800000), top + Emu(80000), Emu(5500000), Emu(420000), [(b, 12, MUTED, False)])

# ---------------------------------------------------------------------------
# SLIDE 12 — Números
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Los números del sprint",
    [("Cantidad, ", MAGENTA, True), ("y sobre todo pruebas", NEAR_BLACK, True)],
    11,
)
nums = [
    ("833", "pruebas del backend en Go"),
    ("322", "pruebas del tablero web"),
    ("58 %", "del código Go es prueba"),
    ("12", "migraciones de base de datos"),
    ("17", "servicios publicados en la API"),
    ("3", "fuentes de reportes integradas"),
    ("16 / 17", "issues del Sprint 3 entregadas"),
    ("18", "compuertas de calidad en verde"),
]
nw = Emu(2500000)
nh = Emu(1100000)
ng = Emu(160000)
for i, (n, lab) in enumerate(nums):
    row, col = divmod(i, 4)
    left = ML + col * (nw + ng)
    top = Emu(2100000) + row * (nh + Emu(150000))
    card(s, left, top, nw, nh, CARD_HOT if i in (0, 1, 6) else CARD)
    add_text(s, left + Emu(100000), top + Emu(200000), nw - Emu(200000), Emu(450000), [(n, 26, MAGENTA, True)], align=PP_ALIGN.CENTER)
    add_text(s, left + Emu(100000), top + Emu(650000), nw - Emu(200000), Emu(350000), [(lab, 11, MUTED, False)], align=PP_ALIGN.CENTER)
add_text(
    s,
    ML,
    Emu(4700000),
    CONTENT_W,
    Emu(300000),
    [("Nada llega a producción sin pasar por:", 13, NEAR_BLACK, True)],
)
add_text(
    s,
    ML,
    Emu(5050000),
    CONTENT_W,
    Emu(900000),
    [
        ("— Frontera de arquitectura · contrato de la API validado · arranque completo con datos de prueba", 12, MUTED, False),
        ("— Lote de 10.000 registros en menos de 5 minutos · reproducibilidad del motor · revisión de estilo", 12, MUTED, False),
    ],
)

# ---------------------------------------------------------------------------
# SLIDE 13 — Demo
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Demostración",
    [("Y ya se puede ver ", MAGENTA, True), ("funcionando", NEAR_BLACK, True)],
    12,
)
s.shapes.add_picture(IMG[8], ML, Emu(2050000), width=Emu(6200000), height=Emu(3800000))
card(s, Emu(7600000), Emu(2050000), Emu(4200000), Emu(1400000), CARD_HOT)
add_text(
    s,
    Emu(7780000),
    Emu(2200000),
    Emu(3840000),
    Emu(300000),
    [("Dirección del entorno desplegado", 13, NEAR_BLACK, True)],
)
add_text(
    s,
    Emu(7780000),
    Emu(2600000),
    Emu(3840000),
    Emu(500000),
    [("main.d2o2d9bqy70o7u.amplifyapp.com", 13, MAGENTA, True)],
)
# QR
s.shapes.add_picture(IMG[9], Emu(10500000), Emu(2200000), width=Emu(900000), height=Emu(900000))
for i, (t, b) in enumerate(
    [
        ("Despliegue automático", "Cada cambio aprobado en main se publica solo: API, base de datos y tablero."),
        ("Cinco roles del reglamento", "Administración, distribución, contabilidad, auditoría y titular ven su parte."),
    ]
):
    top = Emu(3650000) + i * Emu(1200000)
    card(s, Emu(7600000), top, Emu(4200000), Emu(1100000))
    add_text(s, Emu(7780000), top + Emu(180000), Emu(3840000), Emu(300000), [(t, 13, NEAR_BLACK, True)])
    add_text(s, Emu(7780000), top + Emu(500000), Emu(3840000), Emu(450000), [(b, 12, MUTED, False)])

# ---------------------------------------------------------------------------
# SLIDE 14 — Más allá
# ---------------------------------------------------------------------------
s = new_slide()
header_block(
    s,
    "Más allá del sprint",
    [("Por qué ", MAGENTA, True), ("esto importa", NEAR_BLACK, True)],
    13,
)
why = [
    ("Cada cifra, explicable", "De qué recaudo salió, qué reporte la ponderó, qué versión de las reglas se aplicó y qué declaración fijó el porcentaje."),
    ("Auditar es consultar", "La respuesta a una auditoría es una consulta a la bitácora, no una investigación de semanas."),
    ("Para el escritor", "Un portal donde ve el origen de su dinero: qué obra, qué canal, qué periodo y por qué ese monto."),
]
ww = Emu(3428640)
wg = Emu(180000)
for i, (t, b) in enumerate(why):
    left = ML + i * (ww + wg)
    card(s, left, Emu(2100000), ww, Emu(2800000), CARD_HOT if i == 0 else CARD)
    add_text(s, left + Emu(180000), Emu(2350000), ww - Emu(360000), Emu(500000), [(t, 16, NEAR_BLACK, True)])
    add_text(s, left + Emu(180000), Emu(3000000), ww - Emu(360000), Emu(1600000), [(b, 13, MUTED, False)])
card(s, ML, Emu(5100000), CONTENT_W, Emu(900000), CARD_HOT)
add_text(
    s,
    ML + Emu(200000),
    Emu(5300000),
    CONTENT_W - Emu(400000),
    Emu(600000),
    [
        (
            "El siguiente paso del equipo es el proceso completo de liquidación, con las firmas que exige el reglamento antes de que salga un peso.",
            14,
            NEAR_BLACK,
            False,
        )
    ],
)

# ---------------------------------------------------------------------------
# SLIDE 15 — Gracias
# ---------------------------------------------------------------------------
s = new_slide()
upb_header(s)
rect(s, ML, Emu(2200000), Emu(820000), Emu(100080), MAGENTA)
add_text(s, Emu(1700000), Emu(2000000), Emu(5500000), Emu(700000), [("Gracias", 54, MAGENTA, True)])
add_text(
    s,
    ML,
    Emu(2900000),
    Emu(6500000),
    Emu(800000),
    [
        ("Preguntas, y si quieren,", 22, NEAR_BLACK, False),
        ("la demostración en vivo.", 22, NEAR_BLACK, False),
    ],
)
add_text(
    s,
    ML,
    Emu(4000000),
    Emu(6500000),
    Emu(400000),
    [("main.d2o2d9bqy70o7u.amplifyapp.com", 16, MAGENTA, True)],
)
add_text(
    s,
    ML,
    Emu(4600000),
    Emu(6500000),
    Emu(800000),
    [
        ("Intela · REDES SGC · Universidad Pontificia Bolivariana", 13, MUTED, False),
        ("Emanuel Acevedo · Roy Sandoval · Miguel Legarda · Santiago Mendoza", 13, MUTED, False),
    ],
)
s.shapes.add_picture(IMG[1], Emu(7200000), Emu(3000000), width=Emu(4200000), height=Emu(2625000))
upb_footer(s, 14)

prs.save(OUT)
print("saved", OUT)
print("slides", len(prs.slides))
try:
    import shutil
    shutil.copy2(OUT, OUT_DL)
    print("copied", OUT_DL)
except Exception as e:
    print("downloads copy skipped:", e)
